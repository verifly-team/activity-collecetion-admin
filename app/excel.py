"""Excel 导出：三个 sheet —— 报名清单 / 代表商品 / 附件清单。"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ACCENT = "B0703A"

# Excel 会把以这些字符开头的单元格当公式执行（CSV/Excel 注入）
_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")
ACCENT_SOFT = "F7EFE7"
BORDER_COLOR = "E0D9D0"

CATEGORY_CN = {
    "bag": "箱包手袋",
    "watch": "腕表",
    "jewelry": "珠宝配饰",
    "apparel": "服饰",
    "shoe": "鞋履",
    "other": "其他藏品",
}
ZONE_CN = {"A": "A区", "B": "B区", "C": "C区"}
ENTITY_CN = {"company": "公司", "individual": "个体工商户", "personal": "个人"}
KIND_CN = {"photos": "商品照片", "list": "商品清单", "other": "其他"}


def _thin() -> Border:
    side = Side(style="thin", color=BORDER_COLOR)
    return Border(left=side, right=side, top=side, bottom=side)


def _write_sheet(
    ws,
    headers: list[tuple[str, int]],
    rows: list[list[Any]],
    widths: list[int],
) -> None:
    """headers: [(标题, 是否强调)]；widths: 每列宽度。"""
    for idx, (title, _emphasis) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin()
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            # 防公式注入：以 = + - @ 开头的文本强制按字符串存，Excel 不会当公式执行
            if isinstance(value, str) and value[:1] in _FORMULA_PREFIX:
                cell.data_type = "s"
            cell.border = _thin()
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if r_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ACCENT_SOFT)

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def _fmt_time(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        dt = datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone()
        return dt.strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return iso


def _mb(size: int | None) -> float:
    return round((size or 0) / 1024 / 1024, 2)


def build_workbook(
    rows: list[dict[str, Any]],
    files_map: dict[int, list[dict[str, Any]]],
    upload_path=None,
) -> bytes:
    wb = Workbook()

    # ---------------- Sheet 1：报名清单 ----------------
    ws = wb.active
    ws.title = "报名清单"
    headers = [
        ("报名ID", 1),
        ("提交时间", 1),
        ("品牌/店铺", 1),
        ("联系人", 1),
        ("手机", 1),
        ("微信", 1),
        ("邮箱", 1),
        ("城市", 1),
        ("主体类型", 1),
        ("首次参展", 1),
        ("经营品类", 1),
        ("意向区域", 1),
        ("预计带货件数", 1),
        ("价格区间(元)", 1),
        ("二手%", 1),
        ("未使用%", 1),
        ("样衣%", 1),
        ("复刻%", 1),
        ("附件数", 1),
        ("附件大小(MB)", 1),
        ("来源", 1),
    ]
    widths = [8, 17, 22, 10, 14, 14, 22, 9, 12, 10, 18, 12, 13, 16, 8, 9, 8, 8, 8, 13, 8]

    data = []
    for row in rows:
        price = ""
        if row.get("price_min") is not None or row.get("price_max") is not None:
            price = f"{row.get('price_min') or 0} ~ {row.get('price_max') or 0}"
        data.append(
            [
                row.get("id"),
                _fmt_time(row.get("created_at")),
                row.get("brand_name"),
                row.get("contact_name"),
                row.get("phone"),
                row.get("wechat") or "",
                row.get("email") or "",
                row.get("city") or "",
                ENTITY_CN.get(row.get("entity_type") or "", ""),
                "是" if row.get("is_first_time") else "否",
                "、".join(CATEGORY_CN.get(c, c) for c in (row.get("categories") or [])),
                "、".join(ZONE_CN.get(z, z) for z in (row.get("zones") or [])),
                row.get("expected_qty"),
                price,
                row.get("pct_secondhand"),
                row.get("pct_new_stock"),
                row.get("pct_sample"),
                row.get("pct_replica"),
                row.get("file_count") or 0,
                _mb(row.get("file_size")),
                row.get("source") or "",
            ]
        )
    _write_sheet(ws, headers, data, widths)

    # ---------------- Sheet 2：代表商品 ----------------
    ws2 = wb.create_sheet("代表商品")
    headers2 = [
        ("报名ID", 1),
        ("品牌/店铺", 1),
        ("商品及品牌", 1),
        ("年代／系列", 1),
        ("状态及来源", 1),
        ("拟售价(元)", 1),
    ]
    widths2 = [8, 22, 34, 24, 28, 12]
    data2 = []
    for row in rows:
        for item in row.get("items") or []:
            data2.append(
                [
                    row.get("id"),
                    row.get("brand_name"),
                    item.get("name") or "",
                    item.get("era_series") or "",
                    item.get("condition_source") or "",
                    item.get("price"),
                ]
            )
    if not data2:
        data2.append(["", "（本批报名未填写代表商品）", "", "", "", ""])
    _write_sheet(ws2, headers2, data2, widths2)

    # ---------------- Sheet 3：附件清单 ----------------
    ws3 = wb.create_sheet("附件清单")
    headers3 = [
        ("报名ID", 1),
        ("品牌/店铺", 1),
        ("文件名", 1),
        ("类型", 1),
        ("大小(MB)", 1),
        ("上传时间", 1),
        ("包内路径", 1),
        ("状态", 1),
    ]
    widths3 = [8, 22, 30, 12, 11, 17, 46, 10]
    brand_of = {int(r["id"]): r.get("brand_name", "") for r in rows}
    data3 = []
    for app_id, files in files_map.items():
        for f in files:
            exists = ""
            if upload_path is not None:
                exists = "正常" if (upload_path / f["stored_path"]).exists() else "磁盘缺失"
            data3.append(
                [
                    app_id,
                    brand_of.get(app_id, ""),
                    f.get("original_name") or "",
                    KIND_CN.get(f.get("kind") or "", f.get("kind") or ""),
                    _mb(f.get("size")),
                    _fmt_time(f.get("created_at")),
                    f"{app_id:03d}_{brand_of.get(app_id, '')}/附件/{f.get('original_name')}",
                    exists,
                ]
            )
    if not data3:
        data3.append(["", "（本批报名没有上传附件）", "", "", "", "", "", ""])
    _write_sheet(ws3, headers3, data3, widths3)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_filename() -> str:
    return f"报名清单_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"


def zip_folder_name() -> str:
    return f"导出_{datetime.now().strftime('%Y%m%d_%H%M')}"


def zip_filename() -> str:
    return f"报名打包_{datetime.now().strftime('%Y%m%d_%H%M')}.zip"


__all__ = [
    "build_workbook",
    "excel_filename",
    "zip_filename",
    "zip_folder_name",
]
