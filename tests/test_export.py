import io
import os
import sqlite3
import zipfile

import pytest
from fastapi.testclient import TestClient

BASE = "/tmp/vf_export_test"
os.environ.setdefault("DATABASE_URL", f"sqlite:///{BASE}/app.db")
os.environ.setdefault("UPLOAD_DIR", f"{BASE}/uploads")
os.environ.setdefault("EXPORT_TOKEN", "test-token")
os.environ.setdefault("MAX_EXPORT_RECORDS", "2")
os.environ.setdefault("MAX_EXPORT_MB", "50")
os.environ.setdefault("AUDIT_LOG_PATH", f"{BASE}/audit.jsonl")

from app.main import app  # noqa: E402

AUTH = {"Authorization": "Bearer test-token"}


def _seed():
    os.makedirs(f"{BASE}/uploads/1", exist_ok=True)
    os.makedirs(f"{BASE}/uploads/2", exist_ok=True)
    with open(f"{BASE}/uploads/1/a.zip", "wb") as f:
        f.write(b"PK\x03\x04hello")
    conn = sqlite3.connect(f"{BASE}/app.db")
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS application (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          created_at TEXT NOT NULL, brand_name TEXT NOT NULL, contact_name TEXT NOT NULL,
          phone TEXT NOT NULL, wechat TEXT, email TEXT, city TEXT, entity_type TEXT,
          is_first_time INTEGER NOT NULL DEFAULT 0, categories TEXT NOT NULL, zones TEXT NOT NULL,
          expected_qty INTEGER, price_min INTEGER, price_max INTEGER,
          pct_secondhand INTEGER, pct_new_stock INTEGER, pct_sample INTEGER, pct_replica INTEGER,
          items_json TEXT NOT NULL DEFAULT '[]', consent_at TEXT, privacy_version TEXT, source TEXT
        );
        CREATE TABLE IF NOT EXISTS file (
          id TEXT PRIMARY KEY, application_id INTEGER, kind TEXT, original_name TEXT NOT NULL,
          stored_path TEXT NOT NULL, mime TEXT, size INTEGER NOT NULL DEFAULT 0,
          sha256 TEXT, created_at TEXT NOT NULL
        );
        """
    )
    rows = [
        (1, "2026-12-25T06:00:00Z", "拾光中古", "陈小姐", "13800138000", '["bag"]', '["A"]'),
        (2, "2026-12-25T07:00:00Z", "旧物研究所", "李先生", "13900139000", '["watch"]', '["B"]'),
        (3, "2026-12-26T08:00:00Z", "第三家", "王小姐", "13700137000", '["shoe"]', '["C"]'),
    ]
    for r in rows:
        conn.execute(
            "INSERT INTO application (id, created_at, brand_name, contact_name, phone, wechat,"
            " categories, zones, items_json) VALUES (?,?,?,?,?,'wx',?,?,'[]')",
            r,
        )
    conn.execute(
        "INSERT INTO file (id, application_id, original_name, stored_path, size, created_at)"
        " VALUES ('f1', 1, '照片.zip', '1/a.zip', 1024, '2026-12-25T06:00:00Z')"
    )
    conn.commit()
    conn.close()


@pytest.fixture(scope="module")
def client():
    os.makedirs(BASE, exist_ok=True)
    if os.path.exists(f"{BASE}/app.db"):
        os.remove(f"{BASE}/app.db")
    from app import config

    config.settings.database_url = f"sqlite:///{BASE}/app.db"
    config.settings.upload_dir = f"{BASE}/uploads"
    config.settings.export_token = "test-token"
    config.settings.max_records = 2
    config.settings.max_bundle_mb = 50
    _seed()
    with TestClient(app) as c:
        yield c


def test_healthz(client):
    assert client.get("/healthz").json() == {"ok": True}


def test_auth_required(client):
    assert client.get("/api/v1/applications").status_code == 401
    assert client.get("/api/v1/applications", headers=AUTH).status_code == 200


def test_auth_endpoint(client):
    ok = client.post("/api/v1/auth", json={"token": "test-token"})
    assert ok.status_code == 200 and ok.json()["limits"]["max_records"] == 2
    assert client.post("/api/v1/auth", json={"token": "wrong"}).status_code == 401


def test_list_filters(client):
    data = client.get("/api/v1/applications", headers=AUTH).json()
    assert data["total"] == 3
    assert data["items"][0]["id"] == 3  # 倒序

    only_a = client.get("/api/v1/applications?zone=A", headers=AUTH).json()
    assert only_a["total"] == 1

    by_date = client.get("/api/v1/applications?date_from=2026-12-26", headers=AUTH).json()
    assert by_date["total"] == 1

    kw = client.get("/api/v1/applications?keyword=拾光", headers=AUTH).json()
    assert kw["total"] == 1 and kw["items"][0]["file_count"] == 1


def test_excel_export(client):
    from io import BytesIO

    from openpyxl import load_workbook

    res = client.get("/api/v1/applications.xlsx?ids=1,2", headers=AUTH)
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]

    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["报名清单", "代表商品", "附件清单"]

    ws = wb["报名清单"]
    assert ws["A1"].value == "报名ID"
    assert ws.max_row == 3  # 表头 + 2 条
    brands = [ws.cell(row=r, column=3).value for r in (2, 3)]
    assert "拾光中古" in brands

    ws3 = wb["附件清单"]
    assert ws3.max_row >= 2
    assert ws3["A1"].value == "报名ID"


def test_record_limit_enforced(client):
    res = client.get("/api/v1/applications.xlsx?ids=1,2,3", headers=AUTH)
    assert res.status_code == 400
    assert "最多导出 2 条" in res.json()["error"]["message"]


def test_bundle_zip(client):
    res = client.get("/api/v1/bundle.zip?ids=1", headers=AUTH)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/zip"

    zf = zipfile.ZipFile(io.BytesIO(res.content))
    names = zf.namelist()
    assert any("报名清单_" in n and n.endswith(".xlsx") for n in names)
    assert any(n.endswith("照片.zip") for n in names)
    assert any("001_拾光中古/附件" in n for n in names)
    # 不再生成 JSON
    assert not any(n.endswith(".json") for n in names)


def test_bundle_missing_file_listed(client):
    conn = sqlite3.connect(f"{BASE}/app.db")
    conn.execute(
        "INSERT INTO file (id, application_id, original_name, stored_path, size, created_at)"
        " VALUES ('f2', 2, '不存在.zip', '2/none.zip', 10, '2026-12-25T07:00:00Z')"
    )
    conn.commit()
    conn.close()

    res = client.get("/api/v1/bundle.zip?ids=2", headers=AUTH)
    zf = zipfile.ZipFile(io.BytesIO(res.content))
    assert any("_缺失附件清单.txt" in n for n in zf.namelist())


def test_token_via_query_for_download(client):
    res = client.get("/api/v1/applications.xlsx?ids=1&token=test-token")
    assert res.status_code == 200
    assert client.get("/api/v1/applications.xlsx?ids=1&token=bad").status_code == 401


def test_phone_masked_in_list(client):
    """列表接口不应返回完整号码。"""
    data = client.get("/api/v1/applications", headers=AUTH).json()
    phones = [it["phone"] for it in data["items"]]
    assert phones, "应有数据"
    for p in phones:
        assert "****" in p, f"号码未脱敏：{p}"
    assert all("13800138000" not in p for p in phones)


def test_reveal_phone_endpoint(client):
    res = client.get("/api/v1/applications/1/phone", headers=AUTH)
    assert res.status_code == 200
    body = res.json()
    assert body["phone"] == "13800138000"
    assert body["application_id"] == 1


def test_reveal_phone_requires_token(client):
    assert client.get("/api/v1/applications/1/phone").status_code == 401
    assert client.get("/api/v1/applications/999/phone", headers=AUTH).status_code == 404


def test_export_still_has_full_phone(client):
    """导出 Excel 是给业务用的，需要完整号码。"""
    from io import BytesIO

    from openpyxl import load_workbook

    res = client.get("/api/v1/applications.xlsx?ids=1", headers=AUTH)
    wb = load_workbook(BytesIO(res.content))
    ws = wb["报名清单"]
    assert ws.cell(row=2, column=5).value == "13800138000"


# ---------------------------------------------------------------- 安全回归


def test_security_headers(client):
    res = client.get("/healthz")
    assert res.headers["x-content-type-options"] == "nosniff"
    assert res.headers["x-frame-options"] == "DENY"
    assert res.headers["referrer-policy"] == "no-referrer"
    assert "default-src 'self'" in res.headers["content-security-policy"]
    api = client.get("/api/v1/summary", headers=AUTH)
    assert api.headers["cache-control"] == "no-store"


def test_docs_off_by_default(client):
    """/docs、/openapi.json 默认不暴露。"""
    assert client.get("/docs").status_code == 404
    assert client.get("/openapi.json").status_code == 404


def test_excel_formula_injection_neutralised():
    """商家把 =cmd|... 填进品牌名，导出时必须当纯文本存。"""
    from app.excel import build_workbook
    from openpyxl import load_workbook

    row = {
        "id": 1,
        "created_at": "2026-10-29T10:00:00Z",
        "brand_name": "=cmd|' /C calc'!A0",
        "contact_name": "+触类旁通",
        "phone": "13800138000",
        "wechat": "@SUM(1)",
        "email": "",
        "city": "",
        "entity_type": "company",
        "is_first_time": 0,
        "categories": "[]",
        "zones": "[]",
        "expected_qty": 1,
        "items_json": "[]",
    }
    from pathlib import Path

    wb = load_workbook(io.BytesIO(build_workbook([row], {}, Path("/tmp"))))
    ws = wb["报名清单"]
    brand = ws.cell(row=2, column=3)
    assert brand.value == "=cmd|' /C calc'!A0"
    assert brand.data_type == "s", "公式型单元格会被 Excel 执行"
    assert ws.cell(row=2, column=4).data_type == "s"


def test_export_rate_limited(client, monkeypatch):
    from app import main as m

    monkeypatch.setattr(m._export_limiter, "times", 2)
    m._export_limiter._hits.clear()
    ok1 = client.get("/api/v1/applications.xlsx?ids=1", headers=AUTH)
    ok2 = client.get("/api/v1/applications.xlsx?ids=1", headers=AUTH)
    bad = client.get("/api/v1/applications.xlsx?ids=1", headers=AUTH)
    assert ok1.status_code == 200 and ok2.status_code == 200
    assert bad.status_code == 429


# ---------------------------------------------------------------- 并发闸门 / 流式边界


def test_concurrency_gate_basic():
    from app.limits import ConcurrencyGate

    gate = ConcurrencyGate(2, ttl=60)
    a = gate.acquire()
    b = gate.acquire()
    assert a and b and a != b
    assert gate.acquire() is None, "超过上限必须拒绝"
    gate.release(a)
    c = gate.acquire()
    assert c, "释放后应能再拿到名额"
    gate.release(b)
    gate.release(c)
    assert gate.active == 0


def test_concurrency_gate_reaps_stale_lease():
    """客户端中断、生成器没被关闭时，租约超时后必须自动回收。"""
    import time

    from app.limits import ConcurrencyGate

    gate = ConcurrencyGate(1, ttl=0.2)
    lease = gate.acquire()
    assert lease
    assert gate.acquire() is None
    time.sleep(0.35)
    assert gate.acquire(), "超时租约没被回收，功能会被永久锁死"


def test_concurrency_gate_touch_keeps_lease():
    """活着的任务每写一块就续租，不会被误回收。"""
    import time

    from app.limits import ConcurrencyGate

    gate = ConcurrencyGate(1, ttl=0.3)
    lease = gate.acquire()
    for _ in range(4):
        time.sleep(0.15)
        gate.touch(lease)
        assert gate.acquire() is None, "活跃任务被误回收了"
    gate.release(lease)
    assert gate.active == 0


def test_zip_stop_signal_aborts_quickly():
    """收到 stop 后打包线程要及时退出，不能卡在队列上。"""
    import threading
    from pathlib import Path

    from app.export import iter_zip

    stop = threading.Event()
    gen = iter_zip([], {}, Path("/tmp"), stop)
    next(gen)  # 至少能吐第一块
    stop.set()
    for _ in range(50):  # 最多读 50 块
        try:
            next(gen)
        except StopIteration:
            break
    gen.close()


# ---------------------------------------------------------------- 审计日志


def _audit_lines():
    import json as _json

    path = f"{BASE}/audit.jsonl"
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8") as f:
        return [_json.loads(x) for x in f if x.strip()]


def test_export_writes_audit_file(client):
    before = len(_audit_lines())
    client.get("/api/v1/applications.xlsx?ids=1,2", headers=AUTH)
    lines = _audit_lines()
    assert len(lines) > before, "导出没有写审计"
    last = lines[-1]
    assert last["event"] == "export_xlsx"
    assert last["count"] == 2
    assert last["ids"] == "1,2"
    assert last.get("ip")
    assert "ts" in last


def test_reveal_phone_writes_audit(client):
    client.get("/api/v1/applications/1/phone", headers=AUTH)
    hit = [x for x in _audit_lines() if x["event"] == "reveal_phone"]
    assert hit and hit[-1]["application_id"] == 1


def test_audit_api(client):
    client.get("/api/v1/applications.xlsx?ids=1", headers=AUTH)
    res = client.get("/api/v1/audit?limit=10", headers=AUTH)
    assert res.status_code == 200
    body = res.json()
    assert body["ok"] and body["count"] >= 1
    assert all("ts" in i and "event" in i for i in body["items"])
    # 按类型过滤
    only = client.get("/api/v1/audit?event=export_xlsx", headers=AUTH).json()
    assert all(i["event"] == "export_xlsx" for i in only["items"])
    # 需要 token
    assert client.get("/api/v1/audit").status_code == 401


def test_token_can_be_any_length(client, monkeypatch):
    """token 没有长度限制：短的、长的都能用。"""
    from app import auth

    for token in ("x", "short123", "a" * 200):
        monkeypatch.setattr(auth.settings, "export_token", token, raising=False)
        auth._failed.clear()
        assert client.get("/api/v1/summary", headers={"Authorization": f"Bearer {token}"}).status_code == 200
        auth._failed.clear()
        assert client.get("/api/v1/summary", headers={"Authorization": "Bearer wrong"}).status_code == 401
        auth._failed.clear()


def test_token_with_non_ascii_does_not_crash(client, monkeypatch):
    """token 含中文时走 ?token= 应该是 401 而不是 500（compare_digest 的 ASCII 限制）。

    注：HTTP 请求头只能放 ASCII，所以中文 token 只能通过 query 传。
    """
    from app import auth

    monkeypatch.setattr(auth.settings, "export_token", "密钥中文", raising=False)

    auth._failed.clear()
    ok = client.get("/api/v1/summary", params={"token": "密钥中文"})
    auth._failed.clear()
    bad = client.get("/api/v1/summary", params={"token": "密钥中"})
    auth._failed.clear()

    assert ok.status_code == 200
    assert bad.status_code == 401
